# ============================================================
# IMPORTAÇÕES
# ============================================================

from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_from_directory,
    session,
    redirect,
    url_for
)

import atexit
import multiprocessing
import secrets
import os
import queue
import re
import shutil
import sqlite3
import subprocess
import time
import sys
import threading
import traceback
import unicodedata
import uuid
import webbrowser

from datetime import datetime, timedelta
from functools import wraps

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)


# ============================================================
# ESTABILIDADE DO PYTORCH / WHISPER NO WINDOWS
#
# Durante os testes foi identificado que o PyTorch CPU
# encerrava o processo do worker ao iniciar a transcrição
# quando OpenMP/MKL utilizavam múltiplas threads.
#
# Limitando essas bibliotecas a uma thread, o Whisper
# permanece estável dentro do worker separado.
#
# Esta configuração é herdada também pelos subprocessos
# criados pelo multiprocessing no Windows.
# ============================================================

os.environ["OMP_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"


# ============================================================
# MODO DE EXECUÇÃO
#
# O mesmo arquivo funciona:
#
# 1. pelo Python:
#    python app.py
#
# 2. empacotado pelo PyInstaller:
#    Transcrição em Texto.exe
# ============================================================

MODO_EXECUTAVEL = bool(
    getattr(
        sys,
        "frozen",
        False
    )
)


# ============================================================
# PASTA DOS RECURSOS
#
# Em Python:
#     pasta onde está o app.py.
#
# No PyInstaller:
#     sys._MEIPASS contém os arquivos empacotados,
#     incluindo templates e static.
# ============================================================

def obter_pasta_recursos():

    if (
        MODO_EXECUTAVEL
        and hasattr(
            sys,
            "_MEIPASS"
        )
    ):

        return sys._MEIPASS


    return os.path.dirname(
        os.path.abspath(
            __file__
        )
    )


# ============================================================
# PASTA DA APLICAÇÃO
# ============================================================

def obter_pasta_aplicacao():

    if MODO_EXECUTAVEL:

        return os.path.dirname(
            os.path.abspath(
                sys.executable
            )
        )


    return os.path.dirname(
        os.path.abspath(
            __file__
        )
    )


# ============================================================
# PASTA DOCUMENTOS
#
# No EXE não devemos salvar arquivos dentro de Program Files
# ou de outras pastas protegidas.
#
# As transcrições ficarão em:
#
# Documentos\Transcrição em Texto\transcricoes
# ============================================================

def obter_pasta_documentos():

    pasta = os.path.join(
        os.path.expanduser(
            "~"
        ),
        "Documents"
    )


    try:

        os.makedirs(
            pasta,
            exist_ok=True
        )


        return pasta


    except Exception:

        return obter_pasta_aplicacao()


# ============================================================
# PASTAS PRINCIPAIS
# ============================================================

PASTA_RECURSOS = (
    obter_pasta_recursos()
)

PASTA_PROJETO = (
    obter_pasta_aplicacao()
)


# ============================================================
# DADOS DA APLICAÇÃO
#
# Em modo Python:
# mantém o comportamento atual.
#
# Em modo EXE:
# usa a pasta Documentos do usuário.
# ============================================================

if MODO_EXECUTAVEL:

    PASTA_DADOS = os.path.join(
        obter_pasta_documentos(),
        "Transcrição em Texto"
    )

else:

    PASTA_DADOS = (
        PASTA_PROJETO
    )


PASTA_UPLOADS = os.path.join(
    PASTA_DADOS,
    "uploads"
)


PASTA_TRANSCRICOES = os.path.join(
    PASTA_DADOS,
    "transcricoes"
)


CAMINHO_BANCO_HISTORICO = os.path.join(
    PASTA_DADOS,
    "historico.db"
)


os.makedirs(
    PASTA_UPLOADS,
    exist_ok=True
)


os.makedirs(
    PASTA_TRANSCRICOES,
    exist_ok=True
)


# ============================================================
# CACHE DO MODELO WHISPER
#
# No Python normal continuamos usando o cache padrão.
#
# No EXE o modelo ficará em:
#
# AppData\Local\TranscricaoEmTexto\modelos
#
# Assim ele precisa ser baixado apenas na primeira utilização.
# ============================================================

if MODO_EXECUTAVEL:

    pasta_local_appdata = (
        os.environ.get(
            "LOCALAPPDATA",
            os.path.join(
                os.path.expanduser(
                    "~"
                ),
                "AppData",
                "Local"
            )
        )
    )


    PASTA_MODELOS = os.path.join(
        pasta_local_appdata,
        "TranscricaoEmTexto",
        "modelos"
    )


    os.makedirs(
        PASTA_MODELOS,
        exist_ok=True
    )


else:

    PASTA_MODELOS = None


# ============================================================
# TEMPLATES E STATIC
# ============================================================

PASTA_TEMPLATES = os.path.join(
    PASTA_RECURSOS,
    "templates"
)


PASTA_STATIC = os.path.join(
    PASTA_RECURSOS,
    "static"
)


# ============================================================
# CONFIGURAÇÃO FLASK
# ============================================================

app = Flask(

    __name__,

    template_folder=
        PASTA_TEMPLATES,

    static_folder=
        PASTA_STATIC
)


# ============================================================
# SESSÃO E AUTENTICAÇÃO
# ============================================================

CAMINHO_CHAVE_SESSAO = os.path.join(
    PASTA_DADOS,
    "sessao.key"
)


def obter_chave_secreta():

    try:

        if os.path.isfile(
            CAMINHO_CHAVE_SESSAO
        ):

            with open(
                CAMINHO_CHAVE_SESSAO,
                "r",
                encoding="utf-8"
            ) as arquivo:

                chave = (
                    arquivo.read()
                    .strip()
                )


            if chave:

                return chave


        chave = secrets.token_hex(
            32
        )


        with open(
            CAMINHO_CHAVE_SESSAO,
            "w",
            encoding="utf-8"
        ) as arquivo:

            arquivo.write(
                chave
            )


        return chave


    except Exception:

        # Fallback para não impedir a aplicação de abrir.
        # Nesse caso, a sessão será perdida ao reiniciar.
        return secrets.token_hex(
            32
        )


app.config.update(
    SECRET_KEY=
        obter_chave_secreta(),

    SESSION_COOKIE_HTTPONLY=
        True,

    SESSION_COOKIE_SAMESITE=
        "Lax",

    PERMANENT_SESSION_LIFETIME=
        timedelta(
            days=30
        )
)


# ============================================================
# BANCO DE DADOS DO HISTÓRICO
#
# O SQLite é utilizado para manter um histórico local das
# transcrições mesmo depois que a aplicação é fechada.
#
# Em modo EXE:
#
# Documents\Transcrição em Texto\historico.db
#
# Em modo Python:
#
# historico.db na pasta do projeto.
# ============================================================

def conectar_banco_historico():

    conexao = sqlite3.connect(
        CAMINHO_BANCO_HISTORICO,
        timeout=30
    )


    conexao.row_factory = (
        sqlite3.Row
    )


    conexao.execute(
        "PRAGMA busy_timeout = 5000"
    )


    return conexao


# ============================================================
# INICIALIZAR BANCO DO HISTÓRICO
# ============================================================

def inicializar_banco_historico():

    try:

        with conectar_banco_historico() as conexao:

            conexao.execute(
                """
                CREATE TABLE IF NOT EXISTS historico_transcricoes (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    arquivo_original TEXT NOT NULL,
                    tamanho_bytes INTEGER NOT NULL DEFAULT 0,
                    status TEXT NOT NULL,
                    tempo_segundos INTEGER NOT NULL DEFAULT 0,
                    transcricao TEXT,
                    arquivo_docx TEXT,
                    erro TEXT,
                    criado_em TEXT NOT NULL,
                    atualizado_em TEXT NOT NULL
                )
                """
            )


            conexao.execute(
                """
                CREATE TABLE IF NOT EXISTS usuarios (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    nome TEXT NOT NULL,
                    email TEXT NOT NULL,
                    usuario TEXT NOT NULL,
                    senha_hash TEXT NOT NULL,
                    perfil TEXT NOT NULL DEFAULT 'usuario',
                    ativo INTEGER NOT NULL DEFAULT 1,
                    criado_em TEXT NOT NULL,
                    atualizado_em TEXT NOT NULL,
                    ultimo_acesso TEXT
                )
                """
            )


            conexao.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS
                idx_usuarios_usuario_nocase
                ON usuarios (
                    usuario COLLATE NOCASE
                )
                """
            )


            conexao.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS
                idx_usuarios_email_nocase
                ON usuarios (
                    email COLLATE NOCASE
                )
                """
            )


            colunas_historico = {
                registro[
                    "name"
                ]
                for registro
                in conexao.execute(
                    """
                    PRAGMA table_info(
                        historico_transcricoes
                    )
                    """
                ).fetchall()
            }


            if (
                "usuario_id"
                not in colunas_historico
            ):

                conexao.execute(
                    """
                    ALTER TABLE historico_transcricoes
                    ADD COLUMN usuario_id INTEGER
                    """
                )


            conexao.execute(
                """
                CREATE INDEX IF NOT EXISTS
                idx_historico_usuario_id
                ON historico_transcricoes (
                    usuario_id
                )
                """
            )


            admin_legado = conexao.execute(
                """
                SELECT id
                FROM usuarios
                WHERE perfil = 'admin'
                ORDER BY
                    ativo DESC,
                    id ASC
                LIMIT 1
                """
            ).fetchone()


            if (
                admin_legado
                is not None
            ):

                conexao.execute(
                    """
                    UPDATE historico_transcricoes
                    SET usuario_id = ?
                    WHERE usuario_id IS NULL
                    """,
                    (
                        int(
                            admin_legado[
                                "id"
                            ]
                        ),
                    )
                )


            conexao.execute(
                """
                CREATE INDEX IF NOT EXISTS
                idx_historico_criado_em
                ON historico_transcricoes (
                    criado_em DESC
                )
                """
            )


            conexao.execute(
                """
                CREATE INDEX IF NOT EXISTS
                idx_historico_status
                ON historico_transcricoes (
                    status
                )
                """
            )


            agora = datetime.now().isoformat(
                timespec="seconds"
            )


            # Se a aplicação foi encerrada enquanto algum item
            # estava em processamento, esse registro não deve
            # permanecer indefinidamente como "transcrevendo".
            conexao.execute(
                """
                UPDATE historico_transcricoes
                SET
                    status = 'erro',
                    erro = CASE
                        WHEN erro IS NULL OR TRIM(erro) = ''
                        THEN ?
                        ELSE erro
                    END,
                    atualizado_em = ?
                WHERE status = 'transcrevendo'
                """,
                (
                    (
                        "O processamento foi interrompido "
                        "antes de ser concluído."
                    ),
                    agora
                )
            )


            conexao.commit()


        print()
        print(
            "Banco do histórico pronto:"
        )
        print(
            CAMINHO_BANCO_HISTORICO
        )
        print()


        return True


    except Exception as erro:

        print()
        print("!" * 60)
        print(
            "ERRO AO INICIALIZAR O HISTÓRICO"
        )
        print(
            f"Tipo: "
            f"{type(erro).__name__}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print("!" * 60)
        print()


        return False


# ============================================================
# USUÁRIOS
# ============================================================

def contar_usuarios():

    with conectar_banco_historico() as conexao:

        resultado = conexao.execute(
            """
            SELECT COUNT(*) AS total
            FROM usuarios
            """
        ).fetchone()


    return int(
        resultado[
            "total"
        ]
    )


def obter_usuario_por_id(
    usuario_id
):

    if not usuario_id:

        return None


    with conectar_banco_historico() as conexao:

        registro = conexao.execute(
            """
            SELECT
                id,
                nome,
                email,
                usuario,
                senha_hash,
                perfil,
                ativo,
                criado_em,
                atualizado_em,
                ultimo_acesso
            FROM usuarios
            WHERE id = ?
            LIMIT 1
            """,
            (
                int(
                    usuario_id
                ),
            )
        ).fetchone()


    return registro


def obter_usuario_por_login(
    login
):

    login = (
        str(
            login
            or ""
        )
        .strip()
    )


    if not login:

        return None


    with conectar_banco_historico() as conexao:

        registro = conexao.execute(
            """
            SELECT
                id,
                nome,
                email,
                usuario,
                senha_hash,
                perfil,
                ativo,
                criado_em,
                atualizado_em,
                ultimo_acesso
            FROM usuarios
            WHERE
                LOWER(usuario) =
                    LOWER(?)
                OR
                LOWER(email) =
                    LOWER(?)
            LIMIT 1
            """,
            (
                login,
                login
            )
        ).fetchone()


    return registro


def criar_usuario_administrador(
    nome,
    email,
    usuario,
    senha
):

    agora = datetime.now().isoformat(
        timespec="seconds"
    )


    senha_hash = (
        generate_password_hash(
            senha
        )
    )


    with conectar_banco_historico() as conexao:

        cursor = conexao.execute(
            """
            INSERT INTO usuarios (
                nome,
                email,
                usuario,
                senha_hash,
                perfil,
                ativo,
                criado_em,
                atualizado_em,
                ultimo_acesso
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                nome,
                email,
                usuario,
                senha_hash,
                "admin",
                1,
                agora,
                agora,
                agora
            )
        )


        conexao.commit()


    usuario_id = (
        cursor.lastrowid
    )


    associar_historico_legado_ao_usuario(
        usuario_id
    )


    return usuario_id


def associar_historico_legado_ao_usuario(
    usuario_id
):

    if not usuario_id:

        return


    try:

        with conectar_banco_historico() as conexao:

            conexao.execute(
                """
                UPDATE historico_transcricoes
                SET usuario_id = ?
                WHERE usuario_id IS NULL
                """,
                (
                    int(
                        usuario_id
                    ),
                )
            )


            conexao.commit()


    except Exception as erro:

        print()
        print(
            "Não foi possível associar "
            "o histórico legado ao administrador."
        )
        print(
            f"Mensagem: {str(erro)}"
        )
        print()


def atualizar_ultimo_acesso(
    usuario_id
):

    agora = datetime.now().isoformat(
        timespec="seconds"
    )


    with conectar_banco_historico() as conexao:

        conexao.execute(
            """
            UPDATE usuarios
            SET
                ultimo_acesso = ?,
                atualizado_em = ?
            WHERE id = ?
            """,
            (
                agora,
                agora,
                int(
                    usuario_id
                )
            )
        )


        conexao.commit()


def obter_iniciais_usuario(
    nome
):

    partes = [
        parte
        for parte in (
            str(
                nome
                or ""
            )
            .strip()
            .split()
        )
        if parte
    ]


    if not partes:

        return "US"


    if len(
        partes
    ) == 1:

        return (
            partes[
                0
            ][:2]
            .upper()
        )


    return (
        partes[
            0
        ][0]
        +
        partes[
            -1
        ][0]
    ).upper()


def serializar_usuario_interface(
    registro
):

    if registro is None:

        return None


    nome = registro[
        "nome"
    ]


    return {
        "id":
            registro[
                "id"
            ],

        "nome":
            nome,

        "email":
            registro[
                "email"
            ],

        "usuario":
            registro[
                "usuario"
            ],

        "perfil":
            registro[
                "perfil"
            ],

        "ativo":
            bool(
                registro[
                    "ativo"
                ]
            ),

        "iniciais":
            obter_iniciais_usuario(
                nome
            )
    }


def usuario_sessao_atual():

    usuario_id = session.get(
        "usuario_id"
    )


    if not usuario_id:

        return None


    registro = obter_usuario_por_id(
        usuario_id
    )


    if (
        registro is None
        or not bool(
            registro[
                "ativo"
            ]
        )
    ):

        session.clear()


        return None


    return registro


def login_obrigatorio(
    funcao
):

    @wraps(
        funcao
    )
    def wrapper(
        *args,
        **kwargs
    ):

        if (
            usuario_sessao_atual()
            is None
        ):

            return redirect(
                url_for(
                    "login"
                )
            )


        return funcao(
            *args,
            **kwargs
        )


    return wrapper


def login_api_obrigatorio(
    funcao
):

    @wraps(
        funcao
    )
    def wrapper(
        *args,
        **kwargs
    ):

        if (
            usuario_sessao_atual()
            is None
        ):

            return jsonify(
                {
                    "sucesso":
                        False,

                    "autenticacao":
                        False,

                    "mensagem":
                        (
                            "Sua sessão não está ativa. "
                            "Entre novamente no sistema."
                        )
                }
            ), 401


        return funcao(
            *args,
            **kwargs
        )


    return wrapper


# ============================================================
# ACESSO ADMINISTRATIVO
# ============================================================

def admin_api_obrigatorio(
    funcao
):

    @wraps(
        funcao
    )
    def wrapper(
        *args,
        **kwargs
    ):

        registro = (
            usuario_sessao_atual()
        )


        if (
            registro is None
        ):

            return jsonify(
                {
                    "sucesso":
                        False,

                    "autenticacao":
                        False,

                    "mensagem":
                        (
                            "Sua sessão não está ativa. "
                            "Entre novamente no sistema."
                        )
                }
            ), 401


        if (
            registro[
                "perfil"
            ]
            !=
            "admin"
        ):

            return jsonify(
                {
                    "sucesso":
                        False,

                    "mensagem":
                        (
                            "Apenas administradores podem "
                            "gerenciar usuários."
                        )
                }
            ), 403


        return funcao(
            *args,
            **kwargs
        )


    return wrapper


# ============================================================
# CRIAR REGISTRO NO HISTÓRICO
# ============================================================

def criar_registro_historico(
    arquivo_original,
    tamanho_bytes,
    usuario_id
):

    try:

        agora = datetime.now().isoformat(
            timespec="seconds"
        )


        with conectar_banco_historico() as conexao:

            cursor = conexao.execute(
                """
                INSERT INTO historico_transcricoes (
                    arquivo_original,
                    tamanho_bytes,
                    status,
                    tempo_segundos,
                    transcricao,
                    arquivo_docx,
                    erro,
                    criado_em,
                    atualizado_em,
                    usuario_id
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    arquivo_original,
                    int(
                        tamanho_bytes
                        or 0
                    ),
                    "transcrevendo",
                    0,
                    None,
                    None,
                    None,
                    agora,
                    agora,
                    int(
                        usuario_id
                    )
                )
            )


            conexao.commit()


            return cursor.lastrowid


    except Exception as erro:

        print()
        print(
            "Falha ao criar registro "
            "no histórico."
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()


        return None


# ============================================================
# ATUALIZAR REGISTRO DO HISTÓRICO
# ============================================================

def atualizar_registro_historico(
    registro_id,
    status,
    tempo_segundos=0,
    transcricao=None,
    arquivo_docx=None,
    erro=None
):

    if registro_id is None:

        return


    try:

        agora = datetime.now().isoformat(
            timespec="seconds"
        )


        with conectar_banco_historico() as conexao:

            conexao.execute(
                """
                UPDATE historico_transcricoes
                SET
                    status = ?,
                    tempo_segundos = ?,
                    transcricao = ?,
                    arquivo_docx = ?,
                    erro = ?,
                    atualizado_em = ?
                WHERE id = ?
                """,
                (
                    status,
                    max(
                        0,
                        int(
                            tempo_segundos
                            or 0
                        )
                    ),
                    transcricao,
                    arquivo_docx,
                    erro,
                    agora,
                    int(
                        registro_id
                    )
                )
            )


            conexao.commit()


    except Exception as erro_banco:

        print()
        print(
            "Falha ao atualizar registro "
            "do histórico."
        )
        print(
            f"Registro: "
            f"{registro_id}"
        )
        print(
            f"Mensagem: "
            f"{str(erro_banco)}"
        )
        print()


# ============================================================
# LIMPAR HISTÓRICO
# ============================================================

def existe_historico_em_processamento(
    usuario_id
):

    with conectar_banco_historico() as conexao:

        resultado = conexao.execute(
            """
            SELECT COUNT(*) AS total
            FROM historico_transcricoes
            WHERE
                usuario_id = ?
                AND status IN (
                    'aguardando',
                    'transcrevendo',
                    'cancelando'
                )
            """,
            (
                int(
                    usuario_id
                ),
            )
        ).fetchone()


    return (
        int(
            resultado[
                "total"
            ]
        )
        >
        0
    )


def limpar_historico_persistente(
    usuario_id
):

    with conectar_banco_historico() as conexao:

        total = conexao.execute(
            """
            SELECT COUNT(*) AS total
            FROM historico_transcricoes
            WHERE usuario_id = ?
            """,
            (
                int(
                    usuario_id
                ),
            )
        ).fetchone()


        quantidade = int(
            total[
                "total"
            ]
        )


        conexao.execute(
            """
            DELETE FROM historico_transcricoes
            WHERE usuario_id = ?
            """,
            (
                int(
                    usuario_id
                ),
            )
        )


        conexao.commit()


    return quantidade


# ============================================================
# LISTAR HISTÓRICO
# ============================================================

def listar_historico(
    usuario_id
):

    with conectar_banco_historico() as conexao:

        registros = conexao.execute(
            """
            SELECT
                id,
                arquivo_original,
                tamanho_bytes,
                status,
                tempo_segundos,
                transcricao,
                arquivo_docx,
                erro,
                criado_em,
                atualizado_em,
                usuario_id
            FROM historico_transcricoes
            WHERE usuario_id = ?
            ORDER BY id DESC
            """,
            (
                int(
                    usuario_id
                ),
            )
        ).fetchall()


    resultado = []


    for registro in registros:

        nome_docx = (
            registro[
                "arquivo_docx"
            ]
        )


        download_disponivel = False


        if nome_docx:

            caminho_docx = os.path.join(
                PASTA_TRANSCRICOES,
                nome_docx
            )


            download_disponivel = (
                os.path.isfile(
                    caminho_docx
                )
            )


        resultado.append(
            {
                "id":
                    registro[
                        "id"
                    ],

                "arquivo":
                    registro[
                        "arquivo_original"
                    ],

                "tamanho_bytes":
                    registro[
                        "tamanho_bytes"
                    ],

                "status":
                    registro[
                        "status"
                    ],

                "tempo":
                    registro[
                        "tempo_segundos"
                    ],

                "transcricao":
                    (
                        registro[
                            "transcricao"
                        ]
                        or
                        ""
                    ),

                "arquivo_docx":
                    (
                        nome_docx
                        if download_disponivel
                        else None
                    ),

                "download_disponivel":
                    download_disponivel,

                "erro":
                    (
                        registro[
                            "erro"
                        ]
                        or
                        ""
                    ),

                "criado_em":
                    registro[
                        "criado_em"
                    ],

                "atualizado_em":
                    registro[
                        "atualizado_em"
                    ]
            }
        )


    return resultado


# ============================================================
# ESCOPO DO HISTÓRICO POR USUÁRIO
# ============================================================

def contar_historico_usuario(
    usuario_id
):

    with conectar_banco_historico() as conexao:

        resultado = conexao.execute(
            """
            SELECT COUNT(*) AS total
            FROM historico_transcricoes
            WHERE usuario_id = ?
            """,
            (
                int(
                    usuario_id
                ),
            )
        ).fetchone()


    return int(
        resultado[
            "total"
        ]
    )


def usuario_pode_baixar_docx(
    usuario_id,
    nome_arquivo
):

    with conectar_banco_historico() as conexao:

        registro = conexao.execute(
            """
            SELECT id
            FROM historico_transcricoes
            WHERE
                usuario_id = ?
                AND arquivo_docx = ?
                AND status = 'concluido'
            LIMIT 1
            """,
            (
                int(
                    usuario_id
                ),
                nome_arquivo
            )
        ).fetchone()


    return (
        registro
        is not None
    )


# ============================================================
# CONFIGURAR FFMPEG
#
# Procura nesta ordem:
#
# 1. FFmpeg empacotado pelo PyInstaller;
# 2. FFmpeg ao lado do executável;
# 3. FFmpeg instalado no PATH do Windows.
# ============================================================

def configurar_ffmpeg():

    nome_ffmpeg = (
        "ffmpeg.exe"
        if os.name == "nt"
        else
        "ffmpeg"
    )


    pastas_candidatas = [

        os.path.join(
            PASTA_RECURSOS,
            "ffmpeg",
            "bin"
        ),

        os.path.join(
            PASTA_RECURSOS,
            "ffmpeg"
        ),

        PASTA_RECURSOS,

        os.path.join(
            PASTA_PROJETO,
            "ffmpeg",
            "bin"
        ),

        os.path.join(
            PASTA_PROJETO,
            "ffmpeg"
        ),

        PASTA_PROJETO
    ]


    for pasta in pastas_candidatas:

        caminho = os.path.abspath(
            os.path.join(
                pasta,
                nome_ffmpeg
            )
        )


        if not os.path.isfile(
            caminho
        ):

            continue


        pasta_ffmpeg = os.path.dirname(
            caminho
        )


        path_atual = os.environ.get(
            "PATH",
            ""
        )


        if (
            pasta_ffmpeg
            not in path_atual.split(
                os.pathsep
            )
        ):

            os.environ[
                "PATH"
            ] = (
                pasta_ffmpeg
                +
                os.pathsep
                +
                path_atual
            )


        return caminho


    return shutil.which(
        "ffmpeg"
    )


CAMINHO_FFMPEG = (
    configurar_ffmpeg()
)


# ============================================================
# EXTENSÕES PERMITIDAS
# ============================================================

EXTENSOES_PERMITIDAS = {
    ".mp3",
    ".mp4",
    ".wav",
    ".m4a"
}


# ============================================================
# VERSÃO DA APLICAÇÃO
# ============================================================

VERSAO_APLICACAO = "1.2.0"


# ============================================================
# MODELO WHISPER
# ============================================================

MODELO_WHISPER = "small"


# ============================================================
# CONTEXTO TÉCNICO
# ============================================================

CONTEXTO_WHISPER = """
Aula acadêmica em português sobre tecnologia,
ciência de dados, mineração de dados,
análise de dados, técnicas exploratórias,
análise de cluster, clusters, agrupamento de dados,
modelos preditivos e não preditivos,
análise quantitativa e qualitativa,
ponderação arbitrária, atribuição de pesos,
algoritmos, classificação de dados,
tratamento de dados e dados fidedignos.
"""


# ============================================================
# CORREÇÕES SEGURAS
# ============================================================

CORRECOES_SEGURAS = {

    "predictivo":
        "preditivo",

    "predictiva":
        "preditiva",

    "predictivos":
        "preditivos",

    "predictivas":
        "preditivas",

    "filedigno":
        "fidedigno",

    "filedigna":
        "fidedigna",

    "filedignos":
        "fidedignos",

    "filedignas":
        "fidedignas"
}


# ============================================================
# MULTIPROCESSAMENTO
# ============================================================

CONTEXTO_MP = (
    multiprocessing.get_context(
        "spawn"
    )
)


# ============================================================
# CONTROLE DO WORKER
# ============================================================

worker_processo = None

worker_fila_comandos = None

worker_fila_resultados = None


worker_lock = (
    threading.Lock()
)


# ============================================================
# CONTROLE DA TRANSCRIÇÃO
# ============================================================

estado_lock = (
    threading.Lock()
)


processamento_lock = (
    threading.Lock()
)


tarefa_ativa_id = None

evento_cancelamento_ativo = None

# Usuário proprietário da transcrição atualmente em execução.
# Isso impede que outra conta cancele uma tarefa que não iniciou.
tarefa_ativa_usuario_id = None


# ============================================================
# WORKER WHISPER
# ============================================================

def worker_whisper(
    fila_comandos,
    fila_resultados
):

    try:

        print()
        print("=" * 60)
        print("WORKER WHISPER")
        print("Carregando modelo...")
        print(
            f"Modelo: "
            f"{MODELO_WHISPER}"
        )
        print("=" * 60)
        print()


        import whisper


        # ====================================================
        # NO EXE UTILIZA CACHE PRÓPRIO
        # ====================================================

        if (
            MODO_EXECUTAVEL
            and PASTA_MODELOS
        ):

            modelo = (
                whisper.load_model(

                    MODELO_WHISPER,

                    download_root=
                        PASTA_MODELOS
                )
            )


        else:

            modelo = (
                whisper.load_model(
                    MODELO_WHISPER
                )
            )


        print()
        print("=" * 60)
        print(
            "Worker Whisper pronto."
        )
        print(
            f"Modelo carregado: "
            f"{MODELO_WHISPER}"
        )
        print("=" * 60)
        print()


    except Exception as erro:

        print()
        print("!" * 60)
        print(
            "ERRO AO CARREGAR O MODELO WHISPER"
        )
        print(
            f"Tipo: "
            f"{type(erro).__name__}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()

        traceback.print_exc()

        print("!" * 60)
        print()


        return


    # ========================================================
    # LOOP PERSISTENTE
    # ========================================================

    while True:

        try:

            comando = (
                fila_comandos.get()
            )


            if comando is None:

                print()
                print(
                    "Worker Whisper encerrado."
                )
                print()

                break


            tarefa_id = (
                comando.get(
                    "tarefa_id"
                )
            )


            caminho_arquivo = (
                comando.get(
                    "caminho_arquivo"
                )
            )


            nome_arquivo = (
                comando.get(
                    "nome_arquivo"
                )
            )


            print()
            print("=" * 60)
            print(
                f"WORKER - Transcrevendo: "
                f"{nome_arquivo}"
            )
            print(
                f"Tarefa: "
                f"{tarefa_id}"
            )
            print("=" * 60)
            print()


            try:

                print(
                    "Tentativa principal..."
                )

                print()


                resultado = (
                    modelo.transcribe(

                        caminho_arquivo,

                        language="pt",

                        task="transcribe",

                        temperature=0,

                        condition_on_previous_text=True,

                        fp16=False,

                        initial_prompt=
                            CONTEXTO_WHISPER
                    )
                )


                texto = resultado.get(
                    "text",
                    ""
                )


                fila_resultados.put(
                    {
                        "tarefa_id":
                            tarefa_id,

                        "sucesso":
                            True,

                        "texto":
                            texto,

                        "fallback":
                            False
                    }
                )


                print()
                print(
                    f"Worker concluiu: "
                    f"{nome_arquivo}"
                )
                print()


            except RuntimeError as erro:

                mensagem_erro = str(
                    erro
                )


                print()
                print("-" * 60)
                print(
                    "RuntimeError detectado "
                    "no Whisper."
                )
                print(
                    f"Tipo: "
                    f"{type(erro).__name__}"
                )
                print(
                    f"Mensagem: "
                    f"{mensagem_erro}"
                )
                print("-" * 60)
                print()


                if (
                    "cannot reshape tensor of 0 elements"
                    in mensagem_erro.lower()
                ):

                    print()
                    print("!" * 60)
                    print(
                        "Erro de tensor vazio detectado."
                    )
                    print(
                        "Iniciando modo "
                        "de compatibilidade..."
                    )
                    print("!" * 60)
                    print()


                    try:

                        resultado = (
                            modelo.transcribe(

                                caminho_arquivo,

                                language="pt",

                                task="transcribe",

                                temperature=0,

                                condition_on_previous_text=False,

                                fp16=False
                            )
                        )


                        texto = resultado.get(
                            "text",
                            ""
                        )


                        fila_resultados.put(
                            {
                                "tarefa_id":
                                    tarefa_id,

                                "sucesso":
                                    True,

                                "texto":
                                    texto,

                                "fallback":
                                    True
                            }
                        )


                        print()
                        print(
                            "Fallback concluído "
                            "com sucesso."
                        )
                        print()


                    except Exception as erro_fallback:

                        texto_traceback = (
                            traceback.format_exc()
                        )


                        print()
                        print("!" * 60)
                        print(
                            "ERRO NO FALLBACK "
                            "DO WHISPER"
                        )
                        print(
                            f"Tipo: "
                            f"{type(erro_fallback).__name__}"
                        )
                        print(
                            f"Mensagem: "
                            f"{str(erro_fallback)}"
                        )
                        print()
                        print(
                            texto_traceback
                        )
                        print("!" * 60)
                        print()


                        fila_resultados.put(
                            {
                                "tarefa_id":
                                    tarefa_id,

                                "sucesso":
                                    False,

                                "erro":
                                    str(
                                        erro_fallback
                                    ),

                                "tipo_erro":
                                    type(
                                        erro_fallback
                                    ).__name__,

                                "traceback":
                                    texto_traceback
                            }
                        )


                else:

                    texto_traceback = (
                        traceback.format_exc()
                    )


                    print()
                    print(
                        texto_traceback
                    )
                    print()


                    fila_resultados.put(
                        {
                            "tarefa_id":
                                tarefa_id,

                            "sucesso":
                                False,

                            "erro":
                                mensagem_erro,

                            "tipo_erro":
                                type(
                                    erro
                                ).__name__,

                            "traceback":
                                texto_traceback
                        }
                    )


            except Exception as erro:

                texto_traceback = (
                    traceback.format_exc()
                )


                print()
                print("!" * 60)
                print(
                    "ERRO NO WORKER WHISPER"
                )
                print(
                    f"Tipo: "
                    f"{type(erro).__name__}"
                )
                print(
                    f"Mensagem: "
                    f"{str(erro)}"
                )
                print()
                print(
                    texto_traceback
                )
                print("!" * 60)
                print()


                fila_resultados.put(
                    {
                        "tarefa_id":
                            tarefa_id,

                        "sucesso":
                            False,

                        "erro":
                            str(
                                erro
                            ),

                        "tipo_erro":
                            type(
                                erro
                            ).__name__,

                        "traceback":
                            texto_traceback
                    }
                )


        except Exception as erro:

            print()
            print("!" * 60)
            print(
                "ERRO GERAL NO LOOP DO WORKER"
            )
            print(
                f"Tipo: "
                f"{type(erro).__name__}"
            )
            print(
                f"Mensagem: "
                f"{str(erro)}"
            )
            print()

            traceback.print_exc()

            print("!" * 60)
            print()


# ============================================================
# CRIAR WORKER
# ============================================================

def criar_worker_sem_lock():

    global worker_processo
    global worker_fila_comandos
    global worker_fila_resultados


    worker_fila_comandos = (
        CONTEXTO_MP.Queue()
    )


    worker_fila_resultados = (
        CONTEXTO_MP.Queue()
    )


    worker_processo = (
        CONTEXTO_MP.Process(

            target=
                worker_whisper,

            args=(
                worker_fila_comandos,
                worker_fila_resultados
            ),

            daemon=False
        )
    )


    worker_processo.start()


    print()
    print("=" * 60)
    print(
        "Novo processo do Whisper iniciado."
    )
    print(
        f"PID: "
        f"{worker_processo.pid}"
    )
    print("=" * 60)
    print()


# ============================================================
# INICIAR WORKER
# ============================================================

def iniciar_worker():

    with worker_lock:

        if (
            worker_processo is not None
            and worker_processo.is_alive()
        ):

            return


        criar_worker_sem_lock()


# ============================================================
# OBTER WORKER ATIVO
# ============================================================

def obter_worker_ativo():

    global worker_processo
    global worker_fila_comandos
    global worker_fila_resultados


    with worker_lock:

        if (
            worker_processo is None
            or not worker_processo.is_alive()
        ):

            print()
            print(
                "Worker não estava ativo."
            )
            print(
                "Criando novo worker..."
            )
            print()


            criar_worker_sem_lock()


        return (
            worker_processo,
            worker_fila_comandos,
            worker_fila_resultados
        )


# ============================================================
# ENCERRAR PROCESSO DO WORKER
# ============================================================

def encerrar_processo_worker(
    processo
):

    if processo is None:

        return


    if not processo.is_alive():

        return


    pid = processo.pid


    print()
    print(
        f"Encerrando worker PID "
        f"{pid}..."
    )
    print()


    try:

        if os.name == "nt":

            subprocess.run(
                [
                    "taskkill",
                    "/PID",
                    str(
                        pid
                    ),
                    "/T",
                    "/F"
                ],
                stdout=
                    subprocess.DEVNULL,
                stderr=
                    subprocess.DEVNULL,
                check=False
            )


        else:

            processo.terminate()


        processo.join(
            timeout=3
        )


        if processo.is_alive():

            processo.terminate()

            processo.join(
                timeout=2
            )


        if processo.is_alive():

            processo.kill()

            processo.join(
                timeout=2
            )


    except Exception as erro:

        print()
        print(
            "Erro ao encerrar worker:"
        )
        print(
            str(
                erro
            )
        )
        print()


        try:

            processo.terminate()

        except Exception:

            pass


# ============================================================
# REINICIAR WORKER
# ============================================================

def reiniciar_worker():

    global worker_processo
    global worker_fila_comandos
    global worker_fila_resultados


    with worker_lock:

        processo_antigo = (
            worker_processo
        )


        encerrar_processo_worker(
            processo_antigo
        )


        worker_processo = None

        worker_fila_comandos = None

        worker_fila_resultados = None


        print()
        print(
            "Criando novo worker "
            "após cancelamento..."
        )
        print()


        criar_worker_sem_lock()


# ============================================================
# ENCERRAR WORKER AO FECHAR
# ============================================================

def encerrar_worker_final():

    global worker_processo


    try:

        with worker_lock:

            if (
                worker_processo is None
                or not worker_processo.is_alive()
            ):

                return


            try:

                if (
                    worker_fila_comandos
                    is not None
                ):

                    worker_fila_comandos.put(
                        None
                    )


            except Exception:

                pass


            worker_processo.join(
                timeout=2
            )


            if worker_processo.is_alive():

                encerrar_processo_worker(
                    worker_processo
                )


    except Exception:

        pass


# ============================================================
# NORMALIZAR NOME
# ============================================================

def normalizar_nome_arquivo(
    nome_arquivo
):

    if not nome_arquivo:

        return ""


    nome_arquivo = nome_arquivo.replace(
        "\\",
        "/"
    )


    nome_arquivo = os.path.basename(
        nome_arquivo
    )


    return nome_arquivo.strip()


# ============================================================
# EXCLUIR UPLOAD TEMPORÁRIO
# ============================================================

def excluir_upload_temporario(
    caminho_arquivo
):

    if not caminho_arquivo:

        return


    try:

        if os.path.isfile(
            caminho_arquivo
        ):

            os.remove(
                caminho_arquivo
            )


            print()
            print(
                "Upload temporário removido:"
            )
            print(
                caminho_arquivo
            )
            print()


    except PermissionError as erro:

        print()
        print("!" * 60)
        print(
            "Não foi possível remover "
            "o upload temporário."
        )
        print(
            "O arquivo ainda pode estar em uso."
        )
        print(
            f"Caminho: "
            f"{caminho_arquivo}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print("!" * 60)
        print()


    except Exception as erro:

        print()
        print("!" * 60)
        print(
            "Erro ao remover upload temporário."
        )
        print(
            f"Caminho: "
            f"{caminho_arquivo}"
        )
        print(
            f"Tipo: "
            f"{type(erro).__name__}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print("!" * 60)
        print()


# ============================================================
# LIMPAR UPLOADS ÓRFÃOS
# ============================================================

def limpar_uploads_orfaos():

    try:

        nomes = os.listdir(
            PASTA_UPLOADS
        )


        removidos = 0


        for nome in nomes:

            caminho = os.path.join(
                PASTA_UPLOADS,
                nome
            )


            if not os.path.isfile(
                caminho
            ):

                continue


            try:

                os.remove(
                    caminho
                )

                removidos += 1


            except Exception as erro:

                print()
                print(
                    "Não foi possível remover "
                    "um upload órfão:"
                )
                print(
                    caminho
                )
                print(
                    f"Mensagem: "
                    f"{str(erro)}"
                )
                print()


        if removidos > 0:

            print()
            print(
                f"Limpeza inicial: "
                f"{removidos} upload(s) "
                f"temporário(s) removido(s)."
            )
            print()


    except Exception as erro:

        print()
        print(
            "Falha ao verificar uploads "
            "temporários antigos."
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()


# ============================================================
# VALIDAR UPLOAD
# ============================================================

def validar_upload_salvo(
    caminho_arquivo
):

    if not os.path.isfile(
        caminho_arquivo
    ):

        raise RuntimeError(
            "O arquivo enviado não pôde "
            "ser salvo corretamente."
        )


    tamanho = os.path.getsize(
        caminho_arquivo
    )


    if tamanho <= 0:

        raise RuntimeError(
            "O arquivo enviado está vazio."
        )


# ============================================================
# VERIFICAR EXTENSÃO
# ============================================================

def arquivo_permitido(
    nome_arquivo
):

    extensao = os.path.splitext(
        nome_arquivo
    )[1].lower()


    return (
        extensao
        in EXTENSOES_PERMITIDAS
    )


# ============================================================
# MENSAGEM DE ERRO AMIGÁVEL
# ============================================================

def obter_mensagem_erro_amigavel(
    erro
):

    mensagem_original = str(
        erro
    ).strip()


    mensagem = (
        mensagem_original.lower()
    )


    if isinstance(
        erro,
        PermissionError
    ):

        return (
            "Não foi possível acessar ou salvar o arquivo. "
            "Verifique se ele está aberto em outro programa "
            "e tente novamente."
        )


    if (
        "arquivo enviado está vazio"
        in mensagem
    ):

        return (
            "O arquivo enviado está vazio "
            "e não pode ser transcrito."
        )


    if (
        "não pôde ser salvo corretamente"
        in mensagem
    ):

        return (
            "Não foi possível salvar o arquivo "
            "temporariamente para iniciar a transcrição."
        )


    if (
        "whisper não retornou texto"
        in mensagem
    ):

        return (
            "Não foi possível identificar fala "
            "suficiente neste arquivo."
        )


    if (
        "processo do whisper foi encerrado inesperadamente"
        in mensagem
    ):

        return (
            "O mecanismo de transcrição foi "
            "interrompido inesperadamente. "
            "Ele foi reiniciado; tente transcrever "
            "este arquivo novamente."
        )


    if (
        "ffmpeg"
        in mensagem
        and (
            "not found"
            in mensagem
            or "não encontrado"
            in mensagem
            or "no such file"
            in mensagem
        )
    ):

        return (
            "O FFmpeg não foi encontrado. "
            "Verifique a instalação e tente novamente."
        )


    if (
        "invalid data found"
        in mensagem
        or "error opening input"
        in mensagem
        or "could not open"
        in mensagem
        or "failed to load audio"
        in mensagem
    ):

        return (
            "Não foi possível ler o áudio ou vídeo. "
            "O arquivo pode estar corrompido "
            "ou em um formato incompatível."
        )


    if (
        "no space left on device"
        in mensagem
        or "there is not enough space"
        in mensagem
    ):

        return (
            "Não há espaço suficiente em disco "
            "para concluir a operação."
        )


    if (
        "cannot reshape tensor of 0 elements"
        in mensagem
    ):

        return (
            "O Whisper encontrou um erro interno "
            "ao analisar este arquivo. "
            "Tente novamente ou utilize outro arquivo."
        )


    if not mensagem_original:

        return (
            "Ocorreu um erro inesperado "
            "durante o processamento."
        )


    return (
        "Não foi possível concluir a transcrição "
        "deste arquivo. Consulte o terminal "
        "para os detalhes técnicos."
    )


# ============================================================
# LIMPAR TEXTO
# ============================================================

def limpar_texto(
    texto
):

    if not texto:

        return ""


    texto = unicodedata.normalize(
        "NFC",
        texto
    )


    texto = "".join(
        caractere
        for caractere in texto
        if (
            caractere == "\n"
            or caractere == "\t"
            or ord(
                caractere
            ) >= 32
        )
    )


    texto = re.sub(
        r"[ \t]+",
        " ",
        texto
    )


    texto = re.sub(
        r"\s+([,.!?;:])",
        r"\1",
        texto
    )


    texto = re.sub(
        r"([,.!?;:])([^\s])",
        r"\1 \2",
        texto
    )


    texto = re.sub(
        r" {2,}",
        " ",
        texto
    )


    return texto.strip()


# ============================================================
# PRESERVAR CAIXA
# ============================================================

def preservar_caixa(
    original,
    substituicao
):

    if original.isupper():

        return substituicao.upper()


    if original.istitle():

        return substituicao.capitalize()


    return substituicao


# ============================================================
# CORREÇÃO SEGURA
# ============================================================

def corrigir_texto_seguro(
    texto
):

    for (
        errado,
        correto
    ) in CORRECOES_SEGURAS.items():

        padrao = re.compile(
            rf"\b{re.escape(errado)}\b",
            re.IGNORECASE
        )


        texto = padrao.sub(

            lambda correspondencia:
                preservar_caixa(
                    correspondencia.group(
                        0
                    ),
                    correto
                ),

            texto
        )


    return texto


# ============================================================
# SEPARAR FRASES
# ============================================================

def separar_frases(
    texto
):

    texto = texto.strip()


    if not texto:

        return []


    frases = re.split(
        r"(?<=[.!?])\s+",
        texto
    )


    return [
        frase.strip()
        for frase in frases
        if frase.strip()
    ]


# ============================================================
# CRIAR PARÁGRAFOS
# ============================================================

def criar_paragrafos(
    texto,
    frases_por_paragrafo=3
):

    frases = separar_frases(
        texto
    )


    if not frases:

        return texto


    paragrafos = []


    for indice in range(
        0,
        len(
            frases
        ),
        frases_por_paragrafo
    ):

        grupo = frases[
            indice:
            indice
            +
            frases_por_paragrafo
        ]


        paragrafos.append(
            " ".join(
                grupo
            )
        )


    return "\n\n".join(
        paragrafos
    )


# ============================================================
# TRATAR TEXTO
# ============================================================

def tratar_texto(
    texto
):

    texto = limpar_texto(
        texto
    )


    texto = corrigir_texto_seguro(
        texto
    )


    texto = criar_paragrafos(
        texto,
        frases_por_paragrafo=3
    )


    return texto.strip()


# ============================================================
# NÚMERO DE PÁGINA
# ============================================================

def adicionar_numero_pagina(
    paragrafo
):

    run = paragrafo.add_run()


    inicio_campo = OxmlElement(
        "w:fldChar"
    )


    inicio_campo.set(
        qn(
            "w:fldCharType"
        ),
        "begin"
    )


    instrucao = OxmlElement(
        "w:instrText"
    )


    instrucao.set(
        qn(
            "xml:space"
        ),
        "preserve"
    )


    instrucao.text = "PAGE"


    separador = OxmlElement(
        "w:fldChar"
    )


    separador.set(
        qn(
            "w:fldCharType"
        ),
        "separate"
    )


    fim_campo = OxmlElement(
        "w:fldChar"
    )


    fim_campo.set(
        qn(
            "w:fldCharType"
        ),
        "end"
    )


    run._r.append(
        inicio_campo
    )


    run._r.append(
        instrucao
    )


    run._r.append(
        separador
    )


    run._r.append(
        fim_campo
    )


    run.font.name = "Arial"

    run.font.size = Pt(
        8
    )


# ============================================================
# BORDA INFERIOR
# ============================================================

def adicionar_borda_inferior(
    paragrafo
):

    p = paragrafo._p


    pPr = (
        p.get_or_add_pPr()
    )


    pBdr = pPr.find(
        qn(
            "w:pBdr"
        )
    )


    if pBdr is None:

        pBdr = OxmlElement(
            "w:pBdr"
        )


        pPr.append(
            pBdr
        )


    bottom = OxmlElement(
        "w:bottom"
    )


    bottom.set(
        qn(
            "w:val"
        ),
        "single"
    )


    bottom.set(
        qn(
            "w:sz"
        ),
        "6"
    )


    bottom.set(
        qn(
            "w:space"
        ),
        "4"
    )


    bottom.set(
        qn(
            "w:color"
        ),
        "B7B7B7"
    )


    pBdr.append(
        bottom
    )


# ============================================================
# FONTE
# ============================================================

def configurar_run(
    run,
    tamanho,
    negrito=False
):

    run.font.name = "Arial"

    run.font.size = Pt(
        tamanho
    )

    run.bold = negrito


# ============================================================
# NOME ÚNICO PARA WORD
# ============================================================

def gerar_nome_docx_unico(
    nome_sem_extensao
):

    nome_docx = (
        nome_sem_extensao
        +
        ".docx"
    )


    caminho_docx = os.path.join(
        PASTA_TRANSCRICOES,
        nome_docx
    )


    if not os.path.exists(
        caminho_docx
    ):

        return (
            nome_docx,
            caminho_docx
        )


    contador = 2


    while True:

        nome_docx = (
            f"{nome_sem_extensao} "
            f"({contador}).docx"
        )


        caminho_docx = os.path.join(
            PASTA_TRANSCRICOES,
            nome_docx
        )


        if not os.path.exists(
            caminho_docx
        ):

            return (
                nome_docx,
                caminho_docx
            )


        contador += 1


# ============================================================
# CRIAR WORD
# ============================================================

def criar_documento_word(
    texto,
    arquivo_original,
    caminho_saida
):

    documento = Document()


    secao = documento.sections[
        0
    ]


    secao.top_margin = Cm(
        2.5
    )

    secao.bottom_margin = Cm(
        2.5
    )

    secao.left_margin = Cm(
        3
    )

    secao.right_margin = Cm(
        2
    )


    secao.header_distance = Cm(
        1.2
    )

    secao.footer_distance = Cm(
        1.2
    )


    cabecalho = (
        secao.header
    )


    paragrafo_cabecalho = (
        cabecalho.paragraphs[
            0
        ]
    )


    paragrafo_cabecalho.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    run = (
        paragrafo_cabecalho.add_run(
            "Transcrição em Texto"
        )
    )


    configurar_run(
        run,
        9,
        True
    )


    rodape = (
        secao.footer
    )


    paragrafo_rodape = (
        rodape.paragraphs[
            0
        ]
    )


    paragrafo_rodape.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    run = (
        paragrafo_rodape.add_run(
            "Transcrição automática • Página "
        )
    )


    configurar_run(
        run,
        8
    )


    adicionar_numero_pagina(
        paragrafo_rodape
    )


    titulo = (
        documento.add_paragraph()
    )


    titulo.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    titulo.paragraph_format.space_after = Pt(
        8
    )


    run = titulo.add_run(
        "Transcrição em Texto"
    )


    configurar_run(
        run,
        16,
        True
    )


    subtitulo = (
        documento.add_paragraph()
    )


    subtitulo.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    subtitulo.paragraph_format.space_after = Pt(
        12
    )


    run = subtitulo.add_run(
        "Transcrição"
    )


    configurar_run(
        run,
        13,
        True
    )


    nome_sem_extensao = (
        os.path.splitext(
            arquivo_original
        )[0]
    )


    agora = datetime.now()


    data_formatada = (
        agora.strftime(
            "%d/%m/%Y"
        )
    )


    hora_formatada = (
        agora.strftime(
            "%H:%M"
        )
    )


    metadados = (
        documento.add_paragraph()
    )


    metadados.paragraph_format.space_after = Pt(
        10
    )


    run = metadados.add_run(
        f"Arquivo original: "
        f"{nome_sem_extensao}\n"
    )


    configurar_run(
        run,
        9
    )


    run = metadados.add_run(
        f"Modelo de transcrição: "
        f"Whisper {MODELO_WHISPER}\n"
    )


    configurar_run(
        run,
        9
    )


    run = metadados.add_run(
        f"Gerado em: "
        f"{data_formatada} "
        f"às {hora_formatada}"
    )


    configurar_run(
        run,
        9
    )


    adicionar_borda_inferior(
        metadados
    )


    paragrafos = texto.split(
        "\n\n"
    )


    for texto_paragrafo in paragrafos:

        texto_paragrafo = (
            texto_paragrafo.strip()
        )


        if not texto_paragrafo:

            continue


        paragrafo = (
            documento.add_paragraph()
        )


        paragrafo.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )


        formato = (
            paragrafo.paragraph_format
        )


        formato.first_line_indent = Cm(
            1.25
        )


        formato.line_spacing = 1.5


        formato.space_after = Pt(
            8
        )


        run = paragrafo.add_run(
            texto_paragrafo
        )


        configurar_run(
            run,
            11
        )


    documento.save(
        caminho_saida
    )


# ============================================================
# EXECUTAR TRANSCRIÇÃO NO WORKER
# ============================================================

def executar_transcricao_worker(
    caminho_arquivo,
    nome_arquivo,
    usuario_id
):

    global tarefa_ativa_id
    global evento_cancelamento_ativo
    global tarefa_ativa_usuario_id


    (
        processo_local,
        fila_comandos_local,
        fila_resultados_local
    ) = obter_worker_ativo()


    tarefa_id = str(
        uuid.uuid4()
    )


    evento_cancelamento = (
        threading.Event()
    )


    with estado_lock:

        tarefa_ativa_id = (
            tarefa_id
        )


        evento_cancelamento_ativo = (
            evento_cancelamento
        )


        tarefa_ativa_usuario_id = int(
            usuario_id
        )


    fila_comandos_local.put(
        {
            "tarefa_id":
                tarefa_id,

            "caminho_arquivo":
                caminho_arquivo,

            "nome_arquivo":
                nome_arquivo
        }
    )


    print()
    print(
        f"Tarefa enviada ao worker: "
        f"{tarefa_id}"
    )
    print()


    try:

        while True:

            if (
                evento_cancelamento.is_set()
            ):

                print()
                print("=" * 60)
                print(
                    f"Tarefa cancelada: "
                    f"{tarefa_id}"
                )
                print("=" * 60)
                print()


                return {
                    "sucesso":
                        False,

                    "cancelado":
                        True,

                    "erro":
                        "Processamento cancelado pelo usuário."
                }


            if (
                not processo_local.is_alive()
            ):

                if (
                    evento_cancelamento.is_set()
                ):

                    return {
                        "sucesso":
                            False,

                        "cancelado":
                            True,

                        "erro":
                            "Processamento cancelado pelo usuário."
                    }


                print()
                print(
                    "Worker encerrado inesperadamente."
                )
                print(
                    "Tentando preparar "
                    "um novo worker..."
                )
                print()


                try:

                    iniciar_worker()

                except Exception as erro_reinicio:

                    print()
                    print(
                        "Falha ao reiniciar "
                        "o worker:"
                    )
                    print(
                        str(
                            erro_reinicio
                        )
                    )
                    print()


                return {
                    "sucesso":
                        False,

                    "cancelado":
                        False,

                    "erro":
                        "O processo do Whisper foi encerrado inesperadamente."
                }


            try:

                resposta = (
                    fila_resultados_local.get(
                        timeout=0.25
                    )
                )


            except queue.Empty:

                continue


            if (
                resposta.get(
                    "tarefa_id"
                )
                !=
                tarefa_id
            ):

                continue


            return resposta


    finally:

        with estado_lock:

            if (
                tarefa_ativa_id
                ==
                tarefa_id
            ):

                tarefa_ativa_id = None

                evento_cancelamento_ativo = None

                tarefa_ativa_usuario_id = None


# ============================================================
# ADMINISTRAÇÃO DE USUÁRIOS
# ============================================================

def listar_usuarios_admin():

    with conectar_banco_historico() as conexao:

        registros = conexao.execute(
            """
            SELECT
                id,
                nome,
                email,
                usuario,
                perfil,
                ativo,
                criado_em,
                atualizado_em,
                ultimo_acesso
            FROM usuarios
            ORDER BY
                ativo DESC,
                LOWER(nome) ASC,
                id ASC
            """
        ).fetchall()


    return [
        {
            "id":
                registro[
                    "id"
                ],

            "nome":
                registro[
                    "nome"
                ],

            "email":
                registro[
                    "email"
                ],

            "usuario":
                registro[
                    "usuario"
                ],

            "perfil":
                registro[
                    "perfil"
                ],

            "ativo":
                bool(
                    registro[
                        "ativo"
                    ]
                ),

            "criado_em":
                registro[
                    "criado_em"
                ],

            "atualizado_em":
                registro[
                    "atualizado_em"
                ],

            "ultimo_acesso":
                registro[
                    "ultimo_acesso"
                ],

            "iniciais":
                obter_iniciais_usuario(
                    registro[
                        "nome"
                    ]
                )
        }
        for registro
        in registros
    ]


def contar_administradores_ativos():

    with conectar_banco_historico() as conexao:

        resultado = conexao.execute(
            """
            SELECT COUNT(*) AS total
            FROM usuarios
            WHERE
                perfil = 'admin'
                AND ativo = 1
            """
        ).fetchone()


    return int(
        resultado[
            "total"
        ]
    )


def normalizar_dados_usuario(
    dados,
    criando=False
):

    nome = (
        str(
            dados.get(
                "nome",
                ""
            )
        )
        .strip()
    )


    email = (
        str(
            dados.get(
                "email",
                ""
            )
        )
        .strip()
        .lower()
    )


    usuario = (
        str(
            dados.get(
                "usuario",
                ""
            )
        )
        .strip()
    )


    perfil = (
        str(
            dados.get(
                "perfil",
                "usuario"
            )
        )
        .strip()
        .lower()
    )


    ativo_recebido = dados.get(
        "ativo",
        True
    )


    if isinstance(
        ativo_recebido,
        str
    ):

        ativo = (
            ativo_recebido.strip().lower()
            in {
                "1",
                "true",
                "sim",
                "on"
            }
        )

    else:

        ativo = bool(
            ativo_recebido
        )


    senha = str(
        dados.get(
            "senha",
            ""
        )
        or
        ""
    )


    if len(
        nome
    ) < 3:

        raise ValueError(
            "Informe o nome completo do usuário."
        )


    if (
        "@" not in email
        or "." not in email.split(
            "@"
        )[-1]
    ):

        raise ValueError(
            "Informe um e-mail válido."
        )


    if not re.fullmatch(
        r"[A-Za-z0-9._-]{3,30}",
        usuario
    ):

        raise ValueError(
            (
                "O usuário deve ter entre 3 e 30 caracteres "
                "e usar apenas letras, números, ponto, hífen "
                "ou sublinhado."
            )
        )


    if perfil not in {
        "admin",
        "usuario"
    }:

        raise ValueError(
            "Perfil de usuário inválido."
        )


    if (
        criando
        and len(
            senha
        ) < 8
    ):

        raise ValueError(
            "A senha precisa ter pelo menos 8 caracteres."
        )


    if (
        senha
        and len(
            senha
        ) < 8
    ):

        raise ValueError(
            "A nova senha precisa ter pelo menos 8 caracteres."
        )


    return {
        "nome":
            nome,

        "email":
            email,

        "usuario":
            usuario,

        "perfil":
            perfil,

        "ativo":
            ativo,

        "senha":
            senha
    }


def criar_usuario_admin(
    dados
):

    dados = normalizar_dados_usuario(
        dados,
        criando=True
    )


    agora = datetime.now().isoformat(
        timespec="seconds"
    )


    senha_hash = generate_password_hash(
        dados[
            "senha"
        ]
    )


    with conectar_banco_historico() as conexao:

        cursor = conexao.execute(
            """
            INSERT INTO usuarios (
                nome,
                email,
                usuario,
                senha_hash,
                perfil,
                ativo,
                criado_em,
                atualizado_em,
                ultimo_acesso
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                dados[
                    "nome"
                ],
                dados[
                    "email"
                ],
                dados[
                    "usuario"
                ],
                senha_hash,
                dados[
                    "perfil"
                ],
                1
                if dados[
                    "ativo"
                ]
                else 0,
                agora,
                agora,
                None
            )
        )


        conexao.commit()


    return cursor.lastrowid


def atualizar_usuario_admin(
    usuario_id,
    dados,
    usuario_atual_id
):

    registro = obter_usuario_por_id(
        usuario_id
    )


    if registro is None:

        raise LookupError(
            "Usuário não encontrado."
        )


    dados = normalizar_dados_usuario(
        dados,
        criando=False
    )


    proprio_usuario = (
        int(
            usuario_id
        )
        ==
        int(
            usuario_atual_id
        )
    )


    if proprio_usuario:

        if (
            dados[
                "perfil"
            ]
            !=
            registro[
                "perfil"
            ]
        ):

            raise PermissionError(
                "Você não pode alterar o perfil da sua própria conta."
            )


        if (
            dados[
                "ativo"
            ]
            !=
            bool(
                registro[
                    "ativo"
                ]
            )
        ):

            raise PermissionError(
                "Você não pode desativar sua própria conta."
            )


    removendo_admin_ativo = (
        registro[
            "perfil"
        ]
        ==
        "admin"
        and bool(
            registro[
                "ativo"
            ]
        )
        and (
            dados[
                "perfil"
            ]
            !=
            "admin"
            or not dados[
                "ativo"
            ]
        )
    )


    if (
        removendo_admin_ativo
        and contar_administradores_ativos()
        <=
        1
    ):

        raise PermissionError(
            (
                "O sistema precisa manter pelo menos "
                "um Administrador ativo."
            )
        )


    agora = datetime.now().isoformat(
        timespec="seconds"
    )


    with conectar_banco_historico() as conexao:

        if dados[
            "senha"
        ]:

            conexao.execute(
                """
                UPDATE usuarios
                SET
                    nome = ?,
                    email = ?,
                    usuario = ?,
                    senha_hash = ?,
                    perfil = ?,
                    ativo = ?,
                    atualizado_em = ?
                WHERE id = ?
                """,
                (
                    dados[
                        "nome"
                    ],
                    dados[
                        "email"
                    ],
                    dados[
                        "usuario"
                    ],
                    generate_password_hash(
                        dados[
                            "senha"
                        ]
                    ),
                    dados[
                        "perfil"
                    ],
                    1
                    if dados[
                        "ativo"
                    ]
                    else 0,
                    agora,
                    int(
                        usuario_id
                    )
                )
            )

        else:

            conexao.execute(
                """
                UPDATE usuarios
                SET
                    nome = ?,
                    email = ?,
                    usuario = ?,
                    perfil = ?,
                    ativo = ?,
                    atualizado_em = ?
                WHERE id = ?
                """,
                (
                    dados[
                        "nome"
                    ],
                    dados[
                        "email"
                    ],
                    dados[
                        "usuario"
                    ],
                    dados[
                        "perfil"
                    ],
                    1
                    if dados[
                        "ativo"
                    ]
                    else 0,
                    agora,
                    int(
                        usuario_id
                    )
                )
            )


        conexao.commit()


def excluir_usuario_admin(
    usuario_id,
    usuario_atual_id
):

    registro = obter_usuario_por_id(
        usuario_id
    )


    if registro is None:

        raise LookupError(
            "Usuário não encontrado."
        )


    if (
        int(
            usuario_id
        )
        ==
        int(
            usuario_atual_id
        )
    ):

        raise PermissionError(
            "Você não pode excluir sua própria conta."
        )


    if (
        registro[
            "perfil"
        ]
        ==
        "admin"
        and bool(
            registro[
                "ativo"
            ]
        )
        and contar_administradores_ativos()
        <=
        1
    ):

        raise PermissionError(
            (
                "O sistema precisa manter pelo menos "
                "um Administrador ativo."
            )
        )


    quantidade_historico = (
        contar_historico_usuario(
            usuario_id
        )
    )


    if (
        quantidade_historico >
        0
    ):

        raise PermissionError(
            (
                "Este usuário possui "
                f"{quantidade_historico} "
                +
                (
                    "registro no histórico. "
                    if quantidade_historico == 1
                    else "registros no histórico. "
                )
                +
                "Desative a conta em vez de excluí-la."
            )
        )


    with conectar_banco_historico() as conexao:

        conexao.execute(
            """
            DELETE FROM usuarios
            WHERE id = ?
            """,
            (
                int(
                    usuario_id
                ),
            )
        )


        conexao.commit()


# ============================================================
# PRIMEIRO ACESSO
# ============================================================

@app.route(
    "/primeiro-acesso",
    methods=[
        "GET",
        "POST"
    ]
)
def primeiro_acesso():

    if contar_usuarios() > 0:

        return redirect(
            url_for(
                "login"
            )
        )


    erro = ""


    if (
        request.method ==
        "POST"
    ):

        nome = (
            request.form.get(
                "nome",
                ""
            )
            .strip()
        )


        email = (
            request.form.get(
                "email",
                ""
            )
            .strip()
        )


        usuario = (
            request.form.get(
                "usuario",
                ""
            )
            .strip()
        )


        senha = request.form.get(
            "senha",
            ""
        )


        confirmar_senha = (
            request.form.get(
                "confirmar_senha",
                ""
            )
        )


        if len(
            nome
        ) < 3:

            erro = (
                "Informe seu nome completo."
            )


        elif (
            "@" not in email
            or "." not in email.split(
                "@"
            )[-1]
        ):

            erro = (
                "Informe um e-mail válido."
            )


        elif not re.fullmatch(
            r"[A-Za-z0-9._-]{3,30}",
            usuario
        ):

            erro = (
                "O usuário deve ter entre 3 e 30 caracteres "
                "e usar apenas letras, números, ponto, hífen "
                "ou sublinhado."
            )


        elif len(
            senha
        ) < 8:

            erro = (
                "A senha precisa ter pelo menos 8 caracteres."
            )


        elif (
            senha !=
            confirmar_senha
        ):

            erro = (
                "As senhas informadas não coincidem."
            )


        if not erro:

            try:

                usuario_id = (
                    criar_usuario_administrador(
                        nome=
                            nome,

                        email=
                            email,

                        usuario=
                            usuario,

                        senha=
                            senha
                    )
                )


                session.clear()


                session[
                    "usuario_id"
                ] = usuario_id


                session.permanent = (
                    True
                )


                return redirect(
                    url_for(
                        "index"
                    )
                )


            except sqlite3.IntegrityError:

                erro = (
                    "O usuário ou e-mail informado já está cadastrado."
                )


            except Exception as erro_criacao:

                print()
                print(
                    "Erro ao criar administrador:"
                )
                print(
                    str(
                        erro_criacao
                    )
                )
                print()


                erro = (
                    "Não foi possível criar o administrador."
                )


    return render_template(
        "login.html",
        modo=
            "primeiro_acesso",

        erro=
            erro,

        versao=
            VERSAO_APLICACAO
    )


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=[
        "GET",
        "POST"
    ]
)
def login():

    if contar_usuarios() == 0:

        return redirect(
            url_for(
                "primeiro_acesso"
            )
        )


    if (
        usuario_sessao_atual()
        is not None
    ):

        return redirect(
            url_for(
                "index"
            )
        )


    erro = ""


    if (
        request.method ==
        "POST"
    ):

        login_informado = (
            request.form.get(
                "login",
                ""
            )
            .strip()
        )


        senha = request.form.get(
            "senha",
            ""
        )


        manter_conectado = (
            request.form.get(
                "manter_conectado"
            )
            ==
            "1"
        )


        registro = (
            obter_usuario_por_login(
                login_informado
            )
        )


        credenciais_validas = (
            registro is not None
            and bool(
                registro[
                    "ativo"
                ]
            )
            and check_password_hash(
                registro[
                    "senha_hash"
                ],
                senha
            )
        )


        if (
            not credenciais_validas
        ):

            erro = (
                "Usuário/e-mail ou senha inválidos."
            )


        else:

            session.clear()


            session[
                "usuario_id"
            ] = registro[
                "id"
            ]


            session.permanent = (
                manter_conectado
            )


            atualizar_ultimo_acesso(
                registro[
                    "id"
                ]
            )


            return redirect(
                url_for(
                    "index"
                )
            )


    return render_template(
        "login.html",
        modo=
            "login",

        erro=
            erro,

        versao=
            VERSAO_APLICACAO
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route(
    "/logout",
    methods=["POST"]
)
@login_obrigatorio
def logout():

    session.clear()


    return redirect(
        url_for(
            "login"
        )
    )


# ============================================================
# API — USUÁRIOS
# ============================================================

@app.route(
    "/api/usuarios",
    methods=[
        "GET",
        "POST"
    ]
)
@admin_api_obrigatorio
def usuarios_admin():

    try:

        usuario_atual = (
            usuario_sessao_atual()
        )


        if (
            request.method ==
            "GET"
        ):

            registros = (
                listar_usuarios_admin()
            )


            return jsonify(
                {
                    "sucesso":
                        True,

                    "total":
                        len(
                            registros
                        ),

                    "usuario_atual_id":
                        usuario_atual[
                            "id"
                        ],

                    "usuarios":
                        registros
                }
            )


        dados = (
            request.get_json(
                silent=True
            )
            or
            {}
        )


        usuario_id = (
            criar_usuario_admin(
                dados
            )
        )


        return jsonify(
            {
                "sucesso":
                    True,

                "usuario_id":
                    usuario_id,

                "mensagem":
                    "Usuário criado com sucesso."
            }
        ), 201


    except sqlite3.IntegrityError:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Já existe uma conta com esse "
                        "usuário ou e-mail."
                    )
            }
        ), 409


    except ValueError as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    str(
                        erro
                    )
            }
        ), 400


    except Exception as erro:

        print()
        print(
            "ERRO NA ADMINISTRAÇÃO DE USUÁRIOS"
        )
        traceback.print_exc()
        print()


        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Não foi possível concluir "
                        "a operação de usuários."
                    )
            }
        ), 500


@app.route(
    "/api/usuarios/<int:usuario_id>",
    methods=[
        "PUT",
        "DELETE"
    ]
)
@admin_api_obrigatorio
def usuario_admin_item(
    usuario_id
):

    try:

        usuario_atual = (
            usuario_sessao_atual()
        )


        if (
            request.method ==
            "DELETE"
        ):

            excluir_usuario_admin(
                usuario_id=
                    usuario_id,

                usuario_atual_id=
                    usuario_atual[
                        "id"
                    ]
            )


            return jsonify(
                {
                    "sucesso":
                        True,

                    "mensagem":
                        "Usuário excluído com sucesso."
                }
            )


        dados = (
            request.get_json(
                silent=True
            )
            or
            {}
        )


        atualizar_usuario_admin(
            usuario_id=
                usuario_id,

            dados=
                dados,

            usuario_atual_id=
                usuario_atual[
                    "id"
                ]
        )


        return jsonify(
            {
                "sucesso":
                    True,

                "usuario_atualizado_id":
                    usuario_id,

                "recarregar":
                    (
                        int(
                            usuario_id
                        )
                        ==
                        int(
                            usuario_atual[
                                "id"
                            ]
                        )
                    ),

                "mensagem":
                    "Usuário atualizado com sucesso."
            }
        )


    except sqlite3.IntegrityError:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Já existe uma conta com esse "
                        "usuário ou e-mail."
                    )
            }
        ), 409


    except LookupError as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    str(
                        erro
                    )
            }
        ), 404


    except PermissionError as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    str(
                        erro
                    )
            }
        ), 403


    except ValueError as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    str(
                        erro
                    )
            }
        ), 400


    except Exception:

        print()
        print(
            "ERRO AO ALTERAR USUÁRIO"
        )
        traceback.print_exc()
        print()


        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Não foi possível concluir "
                        "a alteração do usuário."
                    )
            }
        ), 500


# ============================================================
# API — DIAGNÓSTICO
# ============================================================

@app.route(
    "/api/status",
    methods=["GET"]
)
@login_api_obrigatorio
def api_status():

    usuario = (
        usuario_sessao_atual()
    )


    return jsonify(
        {
            "sucesso":
                True,

            "api":
                True,

            "versao":
                VERSAO_APLICACAO,

            "usuario_id":
                usuario[
                    "id"
                ],

            "perfil":
                usuario[
                    "perfil"
                ]
        }
    )


# ============================================================
# ROTA PRINCIPAL
# ============================================================

@app.route("/")
@login_obrigatorio
def index():

    usuario = (
        serializar_usuario_interface(
            usuario_sessao_atual()
        )
    )


    return render_template(
        "index.html",
        versao=
            VERSAO_APLICACAO,
        usuario=
            usuario
    )


# ============================================================
# HISTÓRICO PERSISTENTE
# ============================================================

@app.route(
    "/historico",
    methods=[
        "GET",
        "DELETE"
    ]
)
@login_api_obrigatorio
def historico():

    try:

        usuario_atual = (
            usuario_sessao_atual()
        )


        usuario_id = (
            usuario_atual[
                "id"
            ]
        )


        if (
            request.method ==
            "DELETE"
        ):

            if (
                existe_historico_em_processamento(
                    usuario_id
                )
            ):

                return jsonify(
                    {
                        "sucesso":
                            False,

                        "mensagem":
                            (
                                "Existe uma transcrição em andamento. "
                                "Aguarde o processamento terminar "
                                "antes de limpar o histórico."
                            )
                    }
                ), 409


            removidos = (
                limpar_historico_persistente(
                    usuario_id
                )
            )


            return jsonify(
                {
                    "sucesso":
                        True,

                    "removidos":
                        removidos,

                    "mensagem":
                        (
                            "Histórico limpo com sucesso."
                        )
                }
            )


        registros = (
            listar_historico(
                usuario_id
            )
        )


        return jsonify(
            {
                "sucesso":
                    True,

                "total":
                    len(
                        registros
                    ),

                "registros":
                    registros
            }
        )


    except Exception as erro:

        print()
        print("!" * 60)
        print(
            "ERRO AO CONSULTAR O HISTÓRICO"
        )
        print(
            f"Tipo: "
            f"{type(erro).__name__}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()

        traceback.print_exc()

        print("!" * 60)
        print()


        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Não foi possível carregar "
                        "o histórico de transcrições."
                    ),

                "registros":
                    []
            }
        ), 500


# ============================================================
# TRANSCRIÇÃO
# ============================================================

@app.route(
    "/transcrever",
    methods=["POST"]
)
@login_api_obrigatorio
def transcrever():

    usuario_atual = (
        usuario_sessao_atual()
    )


    usuario_id = (
        usuario_atual[
            "id"
        ]
    )


    with processamento_lock:

        try:

            arquivos = (
                request.files.getlist(
                    "arquivos"
                )
            )


            if not arquivos:

                return jsonify(
                    {
                        "sucesso":
                            False,

                        "mensagem":
                            "Nenhum arquivo foi enviado.",

                        "resultados":
                            []
                    }
                ), 400


            resultados = []


            for arquivo in arquivos:

                nome_recebido = (
                    arquivo.filename
                )


                nome_arquivo = (
                    normalizar_nome_arquivo(
                        nome_recebido
                    )
                )


                if not nome_arquivo:

                    resultados.append(
                        {
                            "sucesso":
                                False,

                            "arquivo":
                                "",

                            "erro":
                                "Arquivo sem nome."
                        }
                    )


                    continue


                if not arquivo_permitido(
                    nome_arquivo
                ):

                    resultados.append(
                        {
                            "sucesso":
                                False,

                            "arquivo":
                                nome_arquivo,

                            "erro":
                                "Formato de arquivo não suportado."
                        }
                    )


                    continue


                caminho_arquivo = os.path.join(
                    PASTA_UPLOADS,
                    nome_arquivo
                )


                registro_historico_id = None

                inicio_processamento = None

                tamanho_bytes = 0


                try:

                    print()
                    print("=" * 60)
                    print(
                        f"Arquivo recebido: "
                        f"{nome_recebido}"
                    )
                    print(
                        f"Nome normalizado: "
                        f"{nome_arquivo}"
                    )
                    print(
                        f"Caminho: "
                        f"{caminho_arquivo}"
                    )
                    print("=" * 60)
                    print()


                    arquivo.save(
                        caminho_arquivo
                    )


                    validar_upload_salvo(
                        caminho_arquivo
                    )


                    tamanho_bytes = os.path.getsize(
                        caminho_arquivo
                    )


                    registro_historico_id = (
                        criar_registro_historico(
                            nome_arquivo,
                            tamanho_bytes,
                            usuario_id
                        )
                    )


                    inicio_processamento = (
                        time.monotonic()
                    )


                    print(
                        "Upload salvo e validado com sucesso."
                    )

                    print()


                    resposta_worker = (
                        executar_transcricao_worker(
                            caminho_arquivo,
                            nome_arquivo,
                            usuario_id
                        )
                    )


                    tempo_processamento = int(
                        round(
                            time.monotonic()
                            -
                            inicio_processamento
                        )
                    )


                    if resposta_worker.get(
                        "cancelado",
                        False
                    ):

                        print()
                        print(
                            f"Cancelado: "
                            f"{nome_arquivo}"
                        )
                        print()


                        atualizar_registro_historico(

                            registro_id=
                                registro_historico_id,

                            status=
                                "cancelado",

                            tempo_segundos=
                                tempo_processamento,

                            erro=
                                (
                                    "Processamento cancelado "
                                    "pelo usuário."
                                )
                        )


                        resultados.append(
                            {
                                "sucesso":
                                    False,

                                "cancelado":
                                    True,

                                "arquivo":
                                    nome_arquivo,

                                "historico_id":
                                    registro_historico_id,

                                "erro":
                                    "Processamento cancelado pelo usuário."
                            }
                        )


                        break


                    if not resposta_worker.get(
                        "sucesso",
                        False
                    ):

                        erro_worker = (
                            resposta_worker.get(
                                "erro",
                                "Erro desconhecido no Whisper."
                            )
                        )


                        raise RuntimeError(
                            erro_worker
                        )


                    texto_original = (
                        resposta_worker.get(
                            "texto",
                            ""
                        )
                    )


                    if not texto_original.strip():

                        raise RuntimeError(
                            "O Whisper não retornou texto para este arquivo."
                        )


                    print()
                    print(
                        "Texto recebido do Whisper."
                    )
                    print(
                        f"Caracteres: "
                        f"{len(texto_original)}"
                    )
                    print()


                    texto_tratado = (
                        tratar_texto(
                            texto_original
                        )
                    )


                    print(
                        "Tratamento do texto concluído."
                    )

                    print()


                    nome_sem_extensao = (
                        os.path.splitext(
                            nome_arquivo
                        )[0]
                    )


                    (
                        nome_docx,
                        caminho_docx
                    ) = gerar_nome_docx_unico(
                        nome_sem_extensao
                    )


                    print(
                        f"Gerando Word: "
                        f"{nome_docx}"
                    )

                    print()


                    criar_documento_word(

                        texto=
                            texto_tratado,

                        arquivo_original=
                            nome_arquivo,

                        caminho_saida=
                            caminho_docx
                    )


                    atualizar_registro_historico(

                        registro_id=
                            registro_historico_id,

                        status=
                            "concluido",

                        tempo_segundos=
                            tempo_processamento,

                        transcricao=
                            texto_tratado,

                        arquivo_docx=
                            nome_docx,

                        erro=
                            None
                    )


                    print()
                    print("=" * 60)
                    print(
                        f"Transcrição concluída: "
                        f"{nome_arquivo}"
                    )
                    print(
                        f"Word gerado: "
                        f"{nome_docx}"
                    )
                    print("=" * 60)
                    print()


                    resultados.append(
                        {
                            "sucesso":
                                True,

                            "cancelado":
                                False,

                            "arquivo":
                                nome_arquivo,

                            "historico_id":
                                registro_historico_id,

                            "arquivo_docx":
                                nome_docx,

                            "transcricao":
                                texto_tratado
                        }
                    )


                except Exception as erro:

                    print()
                    print("!" * 60)
                    print(
                        f"ERRO AO PROCESSAR: "
                        f"{nome_arquivo}"
                    )
                    print(
                        f"Tipo: "
                        f"{type(erro).__name__}"
                    )
                    print(
                        f"Mensagem: "
                        f"{str(erro)}"
                    )
                    print()

                    traceback.print_exc()

                    print("!" * 60)
                    print()


                    mensagem_amigavel = (
                        obter_mensagem_erro_amigavel(
                            erro
                        )
                    )


                    tempo_processamento = 0


                    if (
                        inicio_processamento
                        is not None
                    ):

                        tempo_processamento = int(
                            round(
                                time.monotonic()
                                -
                                inicio_processamento
                            )
                        )


                    atualizar_registro_historico(

                        registro_id=
                            registro_historico_id,

                        status=
                            "erro",

                        tempo_segundos=
                            tempo_processamento,

                        erro=
                            mensagem_amigavel
                    )


                    resultados.append(
                        {
                            "sucesso":
                                False,

                            "cancelado":
                                False,

                            "arquivo":
                                nome_arquivo,

                            "historico_id":
                                registro_historico_id,

                            "erro":
                                mensagem_amigavel,

                            "tipo_erro":
                                type(
                                    erro
                                ).__name__
                        }
                    )


                finally:

                    excluir_upload_temporario(
                        caminho_arquivo
                    )


            return jsonify(
                {
                    "sucesso":
                        True,

                    "resultados":
                        resultados
                }
            )


        except Exception as erro:

            print()
            print("#" * 60)
            print(
                "ERRO GERAL NA ROTA /transcrever"
            )
            print(
                f"Tipo: "
                f"{type(erro).__name__}"
            )
            print(
                f"Mensagem: "
                f"{str(erro)}"
            )
            print()

            traceback.print_exc()

            print("#" * 60)
            print()


            mensagem_amigavel = (
                obter_mensagem_erro_amigavel(
                    erro
                )
            )


            return jsonify(
                {
                    "sucesso":
                        False,

                    "mensagem":
                        mensagem_amigavel,

                    "tipo_erro":
                        type(
                            erro
                        ).__name__,

                    "resultados":
                        []
                }
            ), 500


# ============================================================
# CANCELAMENTO
#
# Segurança multiusuário:
# somente a mesma conta que iniciou a transcrição ativa pode
# solicitar o cancelamento.
# ============================================================

@app.route(
    "/cancelar",
    methods=["POST"]
)
@login_api_obrigatorio
def cancelar():

    global tarefa_ativa_id
    global evento_cancelamento_ativo
    global tarefa_ativa_usuario_id


    try:

        usuario_atual = (
            usuario_sessao_atual()
        )


        usuario_atual_id = int(
            usuario_atual[
                "id"
            ]
        )


        # A leitura da tarefa, a validação do proprietário e a
        # sinalização do cancelamento acontecem sob o mesmo lock.
        # Assim evitamos que o estado da tarefa mude no meio da
        # verificação.
        with estado_lock:

            tarefa_id = (
                tarefa_ativa_id
            )


            evento = (
                evento_cancelamento_ativo
            )


            usuario_tarefa_id = (
                tarefa_ativa_usuario_id
            )


            if (
                tarefa_id is None
                or evento is None
            ):

                return jsonify(
                    {
                        "sucesso":
                            True,

                        "cancelado":
                            False,

                        "mensagem":
                            "Não existe uma transcrição ativa."
                    }
                )


            if (
                usuario_tarefa_id
                is None
            ):

                return jsonify(
                    {
                        "sucesso":
                            False,

                        "cancelado":
                            False,

                        "mensagem":
                            (
                                "Não foi possível identificar "
                                "o proprietário da transcrição ativa."
                            )
                    }
                ), 409


            if (
                int(
                    usuario_tarefa_id
                )
                !=
                usuario_atual_id
            ):

                return jsonify(
                    {
                        "sucesso":
                            False,

                        "cancelado":
                            False,

                        "mensagem":
                            (
                                "Você não pode cancelar uma transcrição "
                                "iniciada por outra conta."
                            )
                    }
                ), 403


            # Marca a tarefa correta como cancelada antes de
            # liberar o lock.
            evento.set()


        print()
        print("!" * 60)
        print(
            "CANCELAMENTO IMEDIATO SOLICITADO"
        )
        print(
            f"Tarefa: "
            f"{tarefa_id}"
        )
        print(
            f"Usuário: "
            f"{usuario_atual_id}"
        )
        print("!" * 60)
        print()


        # O worker precisa ser reiniciado porque o Whisper
        # executa de forma bloqueante durante a transcrição.
        reiniciar_worker()


        print()
        print("=" * 60)
        print(
            "Cancelamento concluído."
        )
        print(
            "Novo worker sendo preparado."
        )
        print("=" * 60)
        print()


        return jsonify(
            {
                "sucesso":
                    True,

                "cancelado":
                    True,

                "mensagem":
                    "Transcrição cancelada imediatamente."
            }
        )


    except Exception as erro:

        print()
        print("!" * 60)
        print(
            "ERRO AO CANCELAR TRANSCRIÇÃO"
        )
        print(
            f"Tipo: "
            f"{type(erro).__name__}"
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()

        traceback.print_exc()

        print("!" * 60)
        print()


        return jsonify(
            {
                "sucesso":
                    False,

                "cancelado":
                    False,

                "mensagem":
                    str(
                        erro
                    )
            }
        ), 500



# ============================================================
# DIAGNÓSTICO DO SISTEMA
# ============================================================

def verificar_banco_historico():

    try:

        with conectar_banco_historico() as conexao:

            conexao.execute(
                "SELECT 1"
            ).fetchone()


        return True


    except Exception:

        return False


@app.route(
    "/sistema",
    methods=["GET"]
)
@login_api_obrigatorio
def sistema():

    try:

        whisper_ativo = (
            worker_processo is not None
            and worker_processo.is_alive()
        )


        ffmpeg_ativo = bool(
            CAMINHO_FFMPEG
            and os.path.isfile(
                CAMINHO_FFMPEG
            )
        )


        historico_ativo = (
            verificar_banco_historico()
        )


        return jsonify(
            {
                "sucesso":
                    True,

                "versao":
                    VERSAO_APLICACAO,

                "whisper_ativo":
                    whisper_ativo,

                "ffmpeg_ativo":
                    ffmpeg_ativo,

                "historico_ativo":
                    historico_ativo,

                "modelo":
                    MODELO_WHISPER,

                "idioma":
                    "Português (pt)",

                "dispositivo":
                    "CPU",

                "processamento":
                    "Local • Sequencial",

                "modo_execucao":
                    (
                        "Executável Windows"
                        if MODO_EXECUTAVEL
                        else
                        "Python"
                    ),

                "pasta_transcricoes":
                    PASTA_TRANSCRICOES,

                "caminho_banco":
                    CAMINHO_BANCO_HISTORICO,

                "caminho_ffmpeg":
                    (
                        CAMINHO_FFMPEG
                        or
                        ""
                    )
            }
        )


    except Exception as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Não foi possível obter "
                        "o diagnóstico do sistema."
                    ),

                "erro":
                    str(
                        erro
                    )
            }
        ), 500


# ============================================================
# ABRIR PASTA DE TRANSCRIÇÕES
# ============================================================

@app.route(
    "/abrir-pasta-transcricoes",
    methods=["POST"]
)
@login_api_obrigatorio
def abrir_pasta_transcricoes():

    try:

        os.makedirs(
            PASTA_TRANSCRICOES,
            exist_ok=True
        )


        if os.name == "nt":

            os.startfile(
                PASTA_TRANSCRICOES
            )


        elif sys.platform == "darwin":

            subprocess.Popen(
                [
                    "open",
                    PASTA_TRANSCRICOES
                ]
            )


        else:

            subprocess.Popen(
                [
                    "xdg-open",
                    PASTA_TRANSCRICOES
                ]
            )


        return jsonify(
            {
                "sucesso":
                    True,

                "pasta":
                    PASTA_TRANSCRICOES
            }
        )


    except Exception as erro:

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Não foi possível abrir "
                        "a pasta de transcrições."
                    ),

                "erro":
                    str(
                        erro
                    )
            }
        ), 500


# ============================================================
# DOWNLOAD
# ============================================================

@app.route(
    "/download/<path:nome_arquivo>"
)
@login_obrigatorio
def download(
    nome_arquivo
):

    usuario_atual = (
        usuario_sessao_atual()
    )


    if not usuario_pode_baixar_docx(
        usuario_atual[
            "id"
        ],
        nome_arquivo
    ):

        return (
            "Arquivo não encontrado para esta conta.",
            404
        )


    return send_from_directory(

        PASTA_TRANSCRICOES,

        nome_arquivo,

        as_attachment=True
    )


# ============================================================
# ABRIR NAVEGADOR
# ============================================================

def abrir_navegador_aplicacao():

    try:

        webbrowser.open(
            "http://localhost:5000",
            new=1
        )


    except Exception as erro:

        print()
        print(
            "Não foi possível abrir "
            "o navegador automaticamente."
        )
        print(
            f"Mensagem: "
            f"{str(erro)}"
        )
        print()


# ============================================================
# ERROS DA API
# ============================================================

@app.errorhandler(404)
def tratar_erro_404(
    erro
):

    if request.path.startswith(
        "/api/"
    ):

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Endpoint da API não encontrado: "
                        f"{request.path}"
                    )
            }
        ), 404


    return (
        "Página não encontrada.",
        404
    )


@app.errorhandler(405)
def tratar_erro_405(
    erro
):

    if request.path.startswith(
        "/api/"
    ):

        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "Método HTTP não permitido para "
                        f"{request.path}."
                    )
            }
        ), 405


    return (
        "Método não permitido.",
        405
    )


@app.errorhandler(500)
def tratar_erro_500(
    erro
):

    if request.path.startswith(
        "/api/"
    ):

        print()
        print(
            "ERRO INTERNO NA API"
        )
        print(
            repr(
                erro
            )
        )
        print()


        return jsonify(
            {
                "sucesso":
                    False,

                "mensagem":
                    (
                        "O servidor encontrou um erro interno "
                        "ao processar a solicitação da API."
                    )
            }
        ), 500


    return (
        "Erro interno do servidor.",
        500
    )


# ============================================================
# EXECUTAR
# ============================================================

if __name__ == "__main__":

    multiprocessing.freeze_support()


    print()
    print("=" * 60)
    print(
        "Transcrição em Texto"
    )
    print(
        f"Versão: "
        f"{VERSAO_APLICACAO}"
    )
    print(
        f"Modelo Whisper: "
        f"{MODELO_WHISPER}"
    )
    print(
        "Arquitetura: Worker separado"
    )
    print(
        "Cancelamento imediato: ATIVADO"
    )
    print(
        "Threads OpenMP/MKL: 1"
    )
    print(
        "Modo: "
        +
        (
            "Executável Windows"
            if MODO_EXECUTAVEL
            else
            "Python"
        )
    )
    print("=" * 60)
    print()


    print(
        f"Pasta de recursos: "
        f"{PASTA_RECURSOS}"
    )


    print(
        f"Pasta de dados: "
        f"{PASTA_DADOS}"
    )


    print(
        f"Transcrições: "
        f"{PASTA_TRANSCRICOES}"
    )


    print(
        f"Histórico SQLite: "
        f"{CAMINHO_BANCO_HISTORICO}"
    )


    print()


    if CAMINHO_FFMPEG:

        print(
            "FFmpeg encontrado:"
        )


        print(
            CAMINHO_FFMPEG
        )


    else:

        print(
            "ATENÇÃO: FFmpeg não foi encontrado."
        )


        print(
            "A transcrição não funcionará "
            "até que o FFmpeg esteja disponível."
        )


    print()


    inicializar_banco_historico()


    limpar_uploads_orfaos()


    iniciar_worker()


    atexit.register(
        encerrar_worker_final
    )


    print()
    print("=" * 60)
    print(
        "Servidor iniciado."
    )
    print(
        "Acesse: "
        "http://localhost:5000"
    )
    print("=" * 60)
    print()


    if MODO_EXECUTAVEL:

        temporizador_navegador = (
            threading.Timer(

                1.5,

                abrir_navegador_aplicacao
            )
        )


        temporizador_navegador.daemon = (
            True
        )


        temporizador_navegador.start()


    app.run(

        host=
            "127.0.0.1",

        port=
            5000,

        debug=
            not MODO_EXECUTAVEL,

        use_reloader=
            False,

        threaded=
            True
    )